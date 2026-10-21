from __future__ import annotations

import io
import logging
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)

# Попытка импорта MarkItDown
try:
    from markitdown import MarkItDown
    _markitdown_available = True
    _md = MarkItDown()
except ImportError:
    _markitdown_available = False
    _md = None
    logger.warning("markitdown не установлен, будут использоваться fallback конвертеры")


def convert_docx_to_markdown(file_bytes: bytes) -> str:
    """Извлечение текста и таблиц из DOCX в Markdown формат."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        lines = []

        # Извлекаем параграфы
        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                if p.style.name.startswith("Heading 1"):
                    lines.append(f"# {text}")
                elif p.style.name.startswith("Heading 2"):
                    lines.append(f"## {text}")
                elif p.style.name.startswith("Heading 3"):
                    lines.append(f"### {text}")
                else:
                    lines.append(text)

        # Извлекаем таблицы
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                lines.append(" | ".join(cells))

        return "\n\n".join(lines)
    except Exception as e:
        logger.error("Ошибка fallback парсинга DOCX: %s", e)
        return ""


def convert_pdf_to_markdown(file_bytes: bytes) -> str:
    """Извлечение текста из PDF в Markdown формат."""
    # Попытка pypdf
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        pages_text = []
        for i, page in enumerate(reader.pages):
            txt = page.extract_text() or ""
            clean_lines = [l.strip() for l in txt.splitlines() if l.strip()]
            if clean_lines:
                pages_text.append("\n".join(clean_lines))
        if pages_text:
            return "\n\n".join(pages_text)
    except ImportError:
        pass
    except Exception as e:
        logger.warning("Ошибка pypdf: %s", e)

    # Попытка pdfminer
    try:
        from pdfminer.high_level import extract_text
        text = extract_text(io.BytesIO(file_bytes))
        if text.strip():
            return text.strip()
    except Exception as e:
        logger.error("Ошибка pdfminer: %s", e)

    return ""


def convert_file_to_markdown(file_bytes: bytes, filename: str) -> str:
    """
    Основная функция конвертации входящего файла резюме в Markdown.
    Использует MarkItDown с интеллектуальными fallback-обработчиками.
    """
    ext = Path(filename).suffix.lower()

    # 1. Текстовые файлы
    if ext in [".txt", ".md", ".csv", ".json"]:
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                return file_bytes.decode("cp1251")
            except Exception:
                return file_bytes.decode("utf-8", errors="ignore")

    # 2. Попытка MarkItDown
    if _markitdown_available and _md is not None:
        try:
            # Создаем временный файл для markitdown
            import tempfile
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
                tmp.write(file_bytes)
                tmp_path = tmp.name

            try:
                res = _md.convert(tmp_path)
                if res and res.text_content and len(res.text_content.strip()) > 10:
                    return res.text_content.strip()
            finally:
                Path(tmp_path).unlink(missing_ok=True)
        except Exception as e:
            logger.warning("MarkItDown не смог сконвертировать %s: %s. Пробуем fallback.", filename, e)

    # 3. Fallbacks по типам расширений
    if ext in [".docx", ".doc"]:
        text = convert_docx_to_markdown(file_bytes)
        if text:
            return text

    if ext == ".pdf":
        text = convert_pdf_to_markdown(file_bytes)
        if text:
            return text

    # Последняя надежда — декодировать как текст
    return file_bytes.decode("utf-8", errors="ignore")
