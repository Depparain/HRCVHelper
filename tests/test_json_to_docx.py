from __future__ import annotations

import json
from pathlib import Path
import tempfile
import pytest

from services.json_to_word.json_to_docx import generate_cv
from services.json_to_word import DOCX_TEMPLATE_PATH, STYLE_PATH

BASE_DIR = Path(__file__).resolve().parent.parent
EXAMPLE_JSON_PATH = BASE_DIR / "services" / "json_to_word" / "template" / "example.json"


def test_docx_template_and_style_exist():
    assert DOCX_TEMPLATE_PATH.exists(), f"Шаблон не найден: {DOCX_TEMPLATE_PATH}"
    assert STYLE_PATH.exists(), f"style.json не найден: {STYLE_PATH}"
    assert EXAMPLE_JSON_PATH.exists(), f"example.json не найден: {EXAMPLE_JSON_PATH}"


def test_generate_cv_from_example():
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_output = Path(tmp.name)

    try:
        output_file = generate_cv(
            json_path=EXAMPLE_JSON_PATH,
            template_path=DOCX_TEMPLATE_PATH,
            output_path=tmp_output,
        )

        assert output_file.exists(), "Сгенерированный файл .docx не существует"
        assert output_file.stat().st_size > 50000, "Размер файла подозрительно мал"

        # Проверяем, что файл открывается через python-docx
        from docx import Document
        doc = Document(str(output_file))
        assert len(doc.paragraphs) > 0 or len(doc.tables) > 0, "Документ пуст"
    finally:
        tmp_output.unlink(missing_ok=True)
